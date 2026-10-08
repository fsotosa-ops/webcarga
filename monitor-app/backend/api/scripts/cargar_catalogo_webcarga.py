"""Carga en el catálogo los tipos de documento de la planilla de WebCarga
(HU-C1, entrega 2b, F6).

Fuente: monitor-app/bugs/20261006/Tabla_Resumen_General_IANSA.xlsx, que el
usuario fijó como el ESTÁNDAR WebCarga el 07/10. WebCarga revisa y corrige lo
cargado desde Configuración (decisión del 08/10): no hay planilla de
correspondencias aparte.

- Cada tipo NUEVO se da de alta por el mismo servicio que "Nuevo documento"
  (app/services/catalogo.py): APAGADO, con su alias y auditado. Apagado no se
  le pide a nadie; WebCarga lo activa desde Configuración, con la vista previa.
- Los que YA existen con otro nombre no se tocan (COINCIDENCIAS): cambiarles el
  tipo dejaría vencidos de golpe los ya aprobados, que no tienen fecha de
  emisión. WebCarga los ajusta desde Configuración.
- Las coincidencias dudosas NO se crean (duplicarían un documento) y van a la
  lista de dudas, junto con las filas que la planilla deja ambiguas.

La traducción es LITERAL; lo que no se puede leer literalmente es una duda.

Uso (desde monitor-app/backend/api):
    venv/bin/python -m scripts.cargar_catalogo_webcarga RUTA.xlsx            # simula
    venv/bin/python -m scripts.cargar_catalogo_webcarga RUTA.xlsx --aplicar  # carga
"""
from __future__ import annotations

import argparse
import asyncio
import os
import re
import unicodedata
import urllib.parse
from dataclasses import dataclass
from pathlib import Path
from typing import Optional, Union

from app.schemas.requirement import RequirementCreateBody, VigenciaBody
from app.services.catalogo import codigo_desde_nombre, crear_requisito

# ── Lectura ─────────────────────────────────────────────────────────────────


@dataclass(frozen=True)
class Fila:
    numero: int        # N° dentro de su categoría (columna A)
    categoria: str
    documento: str
    periodicidad: str
    fecha_tope: str
    vigencia: str
    cuando: str


def _texto(valor) -> str:
    return re.sub(r"\s+", " ", str(valor or "")).strip()


def leer_planilla(ruta: Path) -> list[Fila]:
    from openpyxl import load_workbook

    hoja = load_workbook(ruta, read_only=True)["Tabla Resumen General"]
    filas = []
    for a, b, c, d, e, f, _g, h in hoja.iter_rows(min_row=2, max_col=8, values_only=True):
        if not c:
            continue
        filas.append(Fila(int(a), _texto(b), _texto(c).rstrip("."), _texto(d), _texto(e),
                          _texto(f), _texto(h)))
    return filas


# ── Traducción ──────────────────────────────────────────────────────────────


def _clave(texto: str) -> str:
    sin_acentos = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", sin_acentos.lower()).strip()


def _entidad(categoria: str) -> str:
    if categoria.endswith("Empresa"):
        return "CARRIER"
    if categoria.endswith("Trabajador"):
        return "DRIVER"
    return "ASSET"


# El mismo documento que ya está en el catálogo con otro nombre. Se verifica
# en el test contra los códigos vivos.
COINCIDENCIAS = {
    ("CARRIER", "certificado de cumplimiento de obligaciones laborales y previsionales f30 1"): "F30_1",
    ("CARRIER", "certificado afiliacion a mutualidad adherida"): "CERT_MUTUAL",
    ("CARRIER", "comprobante reglamento interno de orden higiene y seguridad riohs timbrado por la autoridad sanitaria y la dt"): "REGLAMENTO_INTERNO",
    ("CARRIER", "procedimientos de trabajo seguro pts"): "PTS_CONTRATISTA",
    ("CARRIER", "politica de seguridad y salud en el trabajo de la empresa contratista"): "POLITICA_SEGURIDAD",
    ("DRIVER", "pasaporte o documento de identidad"): "COPIA_CI_CONDUCTOR",
    ("DRIVER", "contrato de trabajo"): "CONTRATO_TRABAJO",
    ("DRIVER", "registro de entrega de epp"): "ENTREGA_EPP",
    ("DRIVER", "licencia de conducir"): "LICENCIA_CONDUCIR",
    ("DRIVER", "registro de entrega pts"): "PTS_CONDUCTOR",
    ("DRIVER", "hoja de vida del conductor"): "HOJA_DE_VIDA",
    ("DRIVER", "informacion de riesgos laborales irl ex odi ds 44"): "DAS_ODI",
    ("DRIVER", "registro capacitacion programa uso y mantencion epp"): "CAPACITACION_EPP",
    ("DRIVER", "plan de emergencia"): "PLAN_EMERGENCIA",
    ("ASSET", "revision tecnica"): "REVISION_TECNICA",
    ("ASSET", "permiso de circulacion"): "PERMISO_CIRCULACION",
    ("ASSET", "seguro obligatorio soap"): "SOAP",
    ("ASSET", "padron o certificado de inscripcion"): "PADRON",
    ("ASSET", "certificado de emision de gases"): "GASES_CONTAMINANTES",
}

# Probablemente el mismo documento, pero no hay cómo confirmarlo sin WebCarga.
# No se crean: duplicar un documento es peor que dejarlo para después.
COINCIDENCIAS_DUDOSAS = {
    ("CARRIER", "certificado de antecedentes laborales y previsionales f30"): (
        "F30_MULTAS", "¿El F30 de la planilla es nuestro \"F30 Multas\"?"),
    ("CARRIER", "contrato asociado al servicio o cotizacion aceptada y o orden de compra aceptada"): (
        "CONTRATO_WEBCARGA", "¿El \"Contrato asociado al servicio\" es nuestro \"Contrato Webcarga\"?"),
}

# Las especialidades del trabajador (N° dentro de "SSO - Trabajador"): se
# exigen solo cuando el trabajo lo pide, así que van "a pedido". El conductor
# no tiene un atributo de especialidad del que derivarlo.
ESPECIALIDADES = {
    **{n: "trabajo en altura" for n in (19, 20, 21, 22)},
    23: "soldadura",
    **{n: "trabajo en caliente" for n in (24, 25, 26, 27)},
    **{n: "espacios confinados" for n in (28, 29, 30)},
    **{n: "trabajos eléctricos" for n in (31, 32, 33)},
    **{n: "sustancias peligrosas" for n in (34, 35, 36, 37, 38)},
    **{n: "manipulación de alimentos" for n in (39, 40, 41, 42, 43)},
    **{n: "guardias" for n in (44, 46)},
    47: "operador D",
}
_SSO_TRABAJADOR = "Seguridad y Salud Ocupacional (Decreto Supremo 76) - Trabajador"

AVISO_MENSUAL = 5  # Con el aviso general (30), un mensual estaría siempre por vencer.


@dataclass(frozen=True)
class Propuesta:
    fila: Fila
    cuerpo: RequirementCreateBody
    duda: Optional[str] = None


@dataclass(frozen=True)
class Coincidencia:
    fila: Fila
    entidad: str
    codigo: str
    dudosa: bool = False
    duda: Optional[str] = None


def _vigencia(f: Fila) -> tuple[VigenciaBody, Optional[str]]:
    per, vig = _clave(f.periodicidad), _clave(f.vigencia)
    if per == "mensual":
        corte = re.match(r"(\d+)", f.fecha_tope)
        if corte:
            duda = (f'"{f.documento}": la fecha tope dice "{f.fecha_tope}"; se leyó el día {corte.group(1)}.'
                    if not f.fecha_tope.endswith("de cada mes") else None)
            return VigenciaBody(politica="CALENDAR_PERIOD", frequency_months=1,
                                cutoff_day=int(corte.group(1)), period_offset_months=1,
                                warning_days=AVISO_MENSUAL), duda
        return (VigenciaBody(politica="CALENDAR_PERIOD", frequency_months=1, cutoff_day=31,
                             period_offset_months=0, warning_days=AVISO_MENSUAL),
                f'"{f.documento}": mensual sin día tope ("{f.fecha_tope}"); se cargó al último día del mes en curso.')
    if per in ("anual", "bienal"):
        duda = (f'"{f.documento}": dice {f.periodicidad} pero "Vigencia: {f.vigencia}"; se cargó como plazo desde la emisión.'
                if vig != "si" else None)
        return VigenciaBody(politica="ISSUE_PLUS_MONTHS",
                            validity_months=12 if per == "anual" else 24), duda
    if per == "definida por documento":
        return VigenciaBody(politica="REQUIRED"), None
    if per == "no aplica":
        if vig == "si":
            return (VigenciaBody(politica="REQUIRED"),
                    f'"{f.documento}": periodicidad "No aplica" pero "Vigencia: Sí"; se cargó con la fecha del documento.')
        return VigenciaBody(politica="NONE"), None
    raise ValueError(f"Periodicidad desconocida en la fila {f.numero}: {f.periodicidad!r}")


def _exigible(f: Fila, entidad: str) -> tuple[str, Optional[str]]:
    if f.categoria == _SSO_TRABAJADOR and f.numero in ESPECIALIDADES:
        return "ON_REQUEST", None
    cuando = _clave(f.cuando)
    if cuando == "en cuanto se solicita":
        return "ON_REQUEST", None
    if cuando == "al mes siguiente de ingreso trabajador":
        return "MONTH_AFTER_START", None
    if cuando == "al termino del trabajador":
        duda = None
        if not re.search(r"finiquito|traslado", _clave(f.documento)):
            duda = f'"{f.documento}": dice "Al término del trabajador"; parece "al ingreso". Se cargó literal.'
        return "ON_ENTITY_END", duda
    return "ON_ENTITY_START", None


def _nombre(f: Fila) -> str:
    """El de la planilla, salvo las cinco "Entrega de EPP específicos", que se
    llaman igual y se distinguen por su especialidad."""
    if re.search(r"entrega de elementos (de )?proteccion personal especificos", _clave(f.documento)):
        return f"Entrega de EPP específicos ({ESPECIALIDADES[f.numero]})"
    return f.documento


def traducir(f: Fila) -> Union[Propuesta, Coincidencia]:
    entidad = _entidad(f.categoria)
    clave = (entidad, _clave(f.documento))
    if clave in COINCIDENCIAS:
        _, duda = _exigible(f, entidad)
        return Coincidencia(f, entidad, COINCIDENCIAS[clave], duda=duda)
    if clave in COINCIDENCIAS_DUDOSAS:
        codigo, pregunta = COINCIDENCIAS_DUDOSAS[clave]
        return Coincidencia(f, entidad, codigo, dudosa=True, duda=pregunta)
    vigencia, duda_vigencia = _vigencia(f)
    exigible, duda_exigible = _exigible(f, entidad)
    duda = " ".join(d for d in (duda_vigencia, duda_exigible) if d) or None
    cuerpo = RequirementCreateBody(
        name=_nombre(f), target_entity=entidad, requirement_level="LEGAL_MANDATORY",
        vigencia=vigencia, exigible_on=exigible,
    )
    return Propuesta(f, cuerpo, duda)


# ── Ejecución ───────────────────────────────────────────────────────────────


def _url_de_la_base() -> str:
    """El DATABASE_URL de backend/api/.env apunta al host directo, que no
    resuelve desde fuera de Supabase; se reescribe al pooler (aws-1)."""
    from dotenv import dotenv_values

    url = os.environ.get("DATABASE_URL") or dotenv_values(".env").get("DATABASE_URL")
    p = urllib.parse.urlparse(url)
    if "pooler.supabase.com" in (p.hostname or ""):
        return url
    ref = p.hostname.split(".")[1]
    clave = urllib.parse.quote(p.password, safe="")
    return f"postgresql://postgres.{ref}:{clave}@aws-1-us-east-1.pooler.supabase.com:5432/postgres"


def resumen(resultado: list) -> str:
    propuestas = [r for r in resultado if isinstance(r, Propuesta)]
    coincidencias = [r for r in resultado if isinstance(r, Coincidencia)]
    lineas = [f"{len(propuestas)} nuevos · {sum(not c.dudosa for c in coincidencias)} ya existen · "
              f"{sum(c.dudosa for c in coincidencias)} dudosos (no se crean)", ""]
    for p in propuestas:
        v = p.cuerpo.vigencia
        lineas.append(f"  + [{p.cuerpo.target_entity:7}] {codigo_desde_nombre(p.cuerpo.name)[:44]:44} "
                      f"{v.politica:17} {p.cuerpo.exigible_on}")
    lineas += ["", "Ya existen (no se tocan):"]
    lineas += [f"  = {c.codigo:22} ← {c.fila.documento}" for c in coincidencias if not c.dudosa]
    lineas += ["", "Dudas para WebCarga:"]
    lineas += [f"  ? {c.duda}" for c in coincidencias if c.duda]
    lineas += [f"  ? {p.duda}" for p in propuestas if p.duda]
    return "\n".join(lineas)


async def aplicar(conn, resultado: list) -> int:
    """Crea los nuevos en UNA transacción: o entran todos, o ninguno."""
    creados = 0
    async with conn.transaction():
        await conn.execute("SET CONSTRAINTS ALL DEFERRED")
        for r in resultado:
            if isinstance(r, Propuesta):
                await crear_requisito(conn, r.cuerpo, actor=None)
                creados += 1
        await conn.execute("SET CONSTRAINTS ALL IMMEDIATE")
    return creados


async def _main(ruta: Path, aplicar_de_verdad: bool) -> None:
    resultado = [traducir(f) for f in leer_planilla(ruta)]
    print(resumen(resultado))
    if not aplicar_de_verdad:
        print("\nSimulación: no se escribió nada. Con --aplicar se cargan los nuevos, apagados.")
        return
    import asyncpg

    conn = await asyncpg.connect(_url_de_la_base(), statement_cache_size=0)
    try:
        print(f"\nCreados: {await aplicar(conn, resultado)} (apagados).")
    finally:
        await conn.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("planilla", type=Path)
    parser.add_argument("--aplicar", action="store_true")
    argumentos = parser.parse_args()
    asyncio.run(_main(argumentos.planilla, argumentos.aplicar))
